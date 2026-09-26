import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Initialize 16:9 widescreen presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)

# Colors
C_WHITE = RGBColor(255, 255, 255)
C_DARK = RGBColor(17, 24, 39)
C_SLATE = RGBColor(100, 116, 139)
C_NAVY = RGBColor(11, 31, 58)
C_BLUE = RGBColor(9, 132, 227)
C_PURPLE = RGBColor(108, 92, 231)
C_GREEN = RGBColor(0, 184, 148)
C_RED = RGBColor(214, 48, 49)
C_ORANGE = RGBColor(255, 107, 53)
C_LIGHT_BLUE = RGBColor(240, 247, 255)
C_BORDER = RGBColor(203, 213, 225)

# Background fill
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = C_WHITE
bg.line.fill.background()

# ════════════ HEADER ════════════
# 1. Team Pill (Top Left)
pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(0.28), Inches(1.8), Inches(0.55))
pill.fill.solid()
pill.fill.fore_color.rgb = C_WHITE
pill.line.color.rgb = C_PURPLE
pill.line.width = Pt(2.2)
tf = pill.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "NIRMAN"
p.font.name = "Arial"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = C_DARK
p.alignment = PP_ALIGN.CENTER

# 2. Main Title (Center)
title_box = slide.shapes.add_textbox(Inches(2.5), Inches(0.25), Inches(7.5), Inches(0.65))
tf_t = title_box.text_frame
p_t = tf_t.paragraphs[0]
p_t.text = "TECHNICAL APPROACH"
p_t.font.name = "Times New Roman"
p_t.font.size = Pt(32)
p_t.font.bold = True
p_t.font.color.rgb = C_DARK
p_t.alignment = PP_ALIGN.CENTER

# 3. SIH 2026 Logo Box (Top Right)
sih_box = slide.shapes.add_textbox(Inches(10.2), Inches(0.18), Inches(2.6), Inches(0.8))
tf_s = sih_box.text_frame
p_s1 = tf_s.paragraphs[0]
p_s1.text = "SMART INDIA"
p_s1.font.name = "Arial"
p_s1.font.size = Pt(11)
p_s1.font.bold = True
p_s1.font.color.rgb = RGBColor(30, 58, 95)
p_s1.alignment = PP_ALIGN.RIGHT

p_s2 = tf_s.add_paragraph()
p_s2.text = "HACKATHON"
p_s2.font.name = "Arial"
p_s2.font.size = Pt(11)
p_s2.font.bold = True
p_s2.font.color.rgb = RGBColor(30, 58, 95)
p_s2.alignment = PP_ALIGN.RIGHT

p_s3 = tf_s.add_paragraph()
p_s3.text = "2026"
p_s3.font.name = "Arial"
p_s3.font.size = Pt(18)
p_s3.font.bold = True
p_s3.font.color.rgb = RGBColor(30, 58, 95)
p_s3.alignment = PP_ALIGN.RIGHT

# ════════════ SECTION 1: WORKFLOW CONTAINER ════════════
outer_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.05), Inches(12.333), Inches(2.35))
outer_box.fill.solid()
outer_box.fill.fore_color.rgb = C_WHITE
outer_box.line.color.rgb = C_BLUE
outer_box.line.width = Pt(2.0)

# Divider line between pipeline and status
div_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.15), Inches(1.2), Inches(0.025), Inches(2.05))
div_line.fill.solid()
div_line.fill.fore_color.rgb = C_BLUE
div_line.line.fill.background()

# Pipeline Steps (Row 1 & Row 2)
steps = [
    # Row 1
    ("1. Outcome Problem Spec", "Depts post KPI targets & escrow budget instead of specs.", C_PURPLE),
    ("2. GFR 161(iv) Compliance", "Automated engine waives prior turnover & 3-yr track record.", C_BORDER),
    ("3. Milestone Escrow Lock", "Department locks pilot tranches (₹18L) in bank escrow.", C_GREEN),
    ("4. Blind Jury Evaluation", "Multi-expert scoring (Feasibility, Security) with zero bias.", C_RED),
    # Row 2
    ("5. Live IoT Telemetry", "Field sensors stream 15-min telemetry directly into sandbox.", C_BORDER),
    ("6. STQC / IIT Lab Audit", "Independent audit validates telemetry logs & physical KPIs.", C_BLUE),
    ("7. Smart Escrow Release", "Milestone funds release to startup upon verified sign-off.", C_GREEN),
    ("8. National GeM Passport", "1 successful pilot earns credential for pan-India scale-up.", C_BLUE),
]

card_w = Inches(2.05)
card_h = Inches(0.92)
card_x_start = Inches(0.7)
gap_x = Inches(0.38)

for i, (stitle, sdesc, sborder) in enumerate(steps):
    row = i // 4
    col = i % 4
    x = card_x_start + col * (card_w + gap_x)
    y = Inches(1.22) if row == 0 else Inches(2.28)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = sborder
    card.line.width = Pt(1.5)

    tf_c = card.text_frame
    tf_c.word_wrap = True
    p1 = tf_c.paragraphs[0]
    p1.text = stitle
    p1.font.name = "Arial"
    p1.font.size = Pt(8.5)
    p1.font.bold = True
    p1.font.color.rgb = C_DARK
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf_c.add_paragraph()
    p2.text = sdesc
    p2.font.name = "Arial"
    p2.font.size = Pt(7.0)
    p2.font.color.rgb = C_SLATE
    p2.alignment = PP_ALIGN.CENTER

    # Add Arrow if not the last item in row
    if col < 3:
        arr_box = slide.shapes.add_textbox(x + card_w, y + Inches(0.2), gap_x, Inches(0.4))
        p_arr = arr_box.text_frame.paragraphs[0]
        p_arr.text = "➔"
        p_arr.font.name = "Arial"
        p_arr.font.size = Pt(11)
        p_arr.font.bold = True
        p_arr.font.color.rgb = C_BLUE
        p_arr.alignment = PP_ALIGN.CENTER

# Right Status Panel
status_box = slide.shapes.add_textbox(Inches(10.3), Inches(1.15), Inches(2.4), Inches(2.15))
tf_stat = status_box.text_frame
tf_stat.word_wrap = True

p_s_head = tf_stat.paragraphs[0]
p_s_head.text = "PRODUCT STATUS :"
p_s_head.font.name = "Arial"
p_s_head.font.size = Pt(10)
p_s_head.font.bold = True
p_s_head.font.color.rgb = C_NAVY

p_s_hl = tf_stat.add_paragraph()
p_s_hl.text = "90% product build completed and pilot testing is in active progress."
p_s_hl.font.name = "Arial"
p_s_hl.font.size = Pt(8.2)
p_s_hl.font.italic = True
p_s_hl.font.bold = True
p_s_hl.font.color.rgb = C_BLUE

p_s_desc = tf_stat.add_paragraph()
p_s_desc.text = "Interactive department problem studio, GFR 161(iv) Doctor rewrite engine, live telemetry stream simulator, and GeM Passport are fully functional."
p_s_desc.font.name = "Arial"
p_s_desc.font.size = Pt(7.5)
p_s_desc.font.italic = True
p_s_desc.font.color.rgb = RGBColor(51, 65, 85)

p_s_next = tf_stat.add_paragraph()
p_s_next.text = "Next: Ground telemetry pilot with UP Rural Health PHC cold-chain & GeM API gateway."
p_s_next.font.name = "Arial"
p_s_next.font.size = Pt(7.5)
p_s_next.font.color.rgb = C_DARK

# ════════════ SECTION 2: 3 TECH STACK COLUMNS ════════════
col_w = Inches(3.95)
col_h = Inches(3.7)
col_gap = Inches(0.24)
start_x = Inches(0.5)
col_y = Inches(3.52)

tech_data = [
    (
        "🖥️ Frontend Development",
        C_BLUE,
        [
            ("HTML5 / Vanilla JS (ES6+):", "Ultra-fast execution with zero framework overhead and <1.2s page loads."),
            ("Master CSS Grid & Clamping:", "Single responsive adaptive system with touch targets >= 44px on all screens."),
            ("Real-time Telemetry Visualizer:", "SVG dynamic charts streaming 15-min sensor logs & milestone escrow gates."),
            ("Offline Session Engine:", "Persistent client state with instant role switching (Dept / Startup / Validator).")
        ],
        "Tags: HTML5  |  CSS3  |  JS ES6+  |  SVG Canvas  |  LocalStorage"
    ),
    (
        "⚙️ Backend Development",
        RGBColor(217, 119, 6),
        [
            ("Python FastAPI / Node.js:", "High-concurrency async endpoints for challenge broadcasting & bid processing."),
            ("SBI Milestone Escrow Engine:", "Smart release contract simulating tranches locked & disbursed via digital escrow."),
            ("GFR Rule 161(iv) Validator:", "Automated rules parser flagging illegal prior-experience / turnover clauses."),
            ("PostgreSQL & Redis:", "Relational audit trail for jury score normalization with Redis caching for instant queries.")
        ],
        "Tags: Node.js  |  FastAPI  |  PostgreSQL  |  Redis  |  Escrow API"
    ),
    (
        "🔒 GovTech, Security & Trust",
        C_PURPLE,
        [
            ("SHA-256 Telemetry Hashing:", "Tamper-proof hash logs binding raw IoT sensor stream to validator audit records."),
            ("STQC & IIT Verification Protocol:", "Independent third-party test certs adhering to MeitY STQC standards."),
            ("DPIIT & Startup India Integration:", "Direct entity lookup verifying DIPP certificate and founder credentials."),
            ("GeM National Passport Generator:", "Verifiable QR credentials for fast-track L1 procurement exemption.")
        ],
        "Tags: SHA-256  |  STQC Lab  |  DPIIT API  |  DigiLocker  |  GeM Gateway"
    )
]

for idx, (head_text, head_color, bullets, tags_text) in enumerate(tech_data):
    cx = start_x + idx * (col_w + col_gap)

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, col_y, col_w, col_h)
    box.fill.solid()
    box.fill.fore_color.rgb = C_WHITE
    box.line.color.rgb = C_BORDER
    box.line.width = Pt(1.5)

    tf_b = box.text_frame
    tf_b.word_wrap = True

    # Title
    p_head = tf_b.paragraphs[0]
    p_head.text = head_text
    p_head.font.name = "Arial"
    p_head.font.size = Pt(13)
    p_head.font.bold = True
    p_head.font.color.rgb = head_color

    # Bullets
    for b_title, b_desc in bullets:
        pb = tf_b.add_paragraph()
        pb.text = f"• {b_title} {b_desc}"
        pb.font.name = "Arial"
        pb.font.size = Pt(8.5)
        pb.font.color.rgb = RGBColor(71, 85, 105)

    # Tags at bottom
    pt = tf_b.add_paragraph()
    pt.text = tags_text
    pt.font.name = "Arial"
    pt.font.size = Pt(8.0)
    pt.font.bold = True
    pt.font.color.rgb = C_NAVY

# ════════════ BOTTOM ACCENT BAR ════════════
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.35), Inches(13.333), Inches(0.15))
bar.fill.solid()
bar.fill.fore_color.rgb = C_BLUE
bar.line.fill.background()

output_path = "/Users/nishant/Desktop/NIRMAN/NIRMAN_Technical_Approach_Slide.pptx"
prs.save(output_path)
print(f"Successfully generated: {output_path}")
